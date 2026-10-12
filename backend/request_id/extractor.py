
import os
import glob
import logging

import openpyxl

import config

logger = logging.getLogger(__name__)

SHEET_NAME = 'Contract Request'
KEY_COLUMN = config.MATCH_KEY_COLUMN     # 'Record Reference Id'
TARGET_COLUMNS = config.TARGET_COLUMNS   # ['Request ID', 'Record ID', 'File Names']

# The site's Record Versions table only exposes the full vault folder
# path (e.g. "MGN INDIA ATTACHMENTS/DMIN10030 - SOlus/Mashreq Agreement
# Solus.pdf"), but the spec's own File Names example is a plain filename
# ("Ebix NDA - Bilateral (1) (1).pdf"). Rather than pick one and discard
# the other, we keep File Names as the plain filename and add this
# column -- auto-inserted immediately after File Names if the workbook
# doesn't already have it -- to hold the full path too.
FULL_PATH_COLUMN = 'File Path'
ALL_TARGET_COLUMNS = TARGET_COLUMNS + [FULL_PATH_COLUMN]

# The Supporting Document sheet is the uploads team's own per-document
# tracking tab, pre-populated with one row per expected supporting file
# (Record Reference Id + Document Reference Id + the intended filename),
# confirmed against a real export (MSA_2302.xlsx). We match scraped
# per-document rows back onto it by (Record Reference Id, Document
# Reference Id) and fill in these columns -- created automatically if
# missing, same as the Contract Request columns above.
SUPPORTING_SHEET_NAME = 'Supporting Document'
SUPPORTING_KEY_COLUMNS = ('Record Reference Id', 'Document Reference Id')
SUPPORTING_TARGET_COLUMNS = [
    'Supporting Document ID', 'Filename', 'File Path', 'Status',
    'Error Message', 'Parent Request ID', 'Parent Record ID',
]


def find_input_workbook(input_dir=None):
    """Return the path to the single .xlsx file sitting in input/, or None.
    Ignores Excel's own lock files (~$...)."""
    input_dir = input_dir or config.INPUT_DIR
    candidates = [
        f for f in glob.glob(os.path.join(input_dir, '*.xlsx'))
        if not os.path.basename(f).startswith('~$')
    ]
    if not candidates:
        return None
    if len(candidates) > 1:
        logger.warning(f'multiple .xlsx files in {input_dir}, using the most recently modified one')
        candidates.sort(key=os.path.getmtime, reverse=True)
    return candidates[0]


def _header_map(ws):
    """Column name -> 1-based column index, read from row 1. Never
    hardcode column letters -- different application-type templates
    (MSA/NDA/SOW/...) can order/include custom fields differently."""
    headers = {}
    for cell in ws[1]:
        if cell.value:
            headers[str(cell.value).strip()] = cell.column
    missing = [c for c in [KEY_COLUMN] + TARGET_COLUMNS if c not in headers]
    if missing:
        raise ValueError(f'sheet {ws.title!r} is missing expected column(s): {missing}')
    return headers


def load_workbook(path):
    return openpyxl.load_workbook(path)


def _is_blank(value):
    return value is None or str(value).strip() == ''


# Visual order the uploads team expects these columns in, immediately
# after the Record Reference Id key column -- confirmed against the
# team's own reference screenshot (Record ID | Request ID | File Names |
# File Path). Note this differs from TARGET_COLUMNS' order, which is
# just a validation list and unrelated to layout.
BASE_INSERT_ORDER = ['Record ID', 'Request ID', 'File Names']


def ensure_output_columns(wb, path):
    """Idempotently create every column the pipeline writes to on the
    Contract Request sheet: Record ID / Request ID / File Names / File
    Path, inserted together right after 'Record Reference Id' (not
    appended past every custom field at the far right, which is easy to
    miss without scrolling -- confirmed live against a fresh uploads-team
    export, MSA_2302.xlsx, that had none of the four yet). A workbook
    from a previous (already-upgraded) run is left alone -- if 'File
    Names' already exists somewhere else but 'File Path' doesn't, File
    Path is inserted right after wherever File Names already is, rather
    than forcing a relocation. Saves immediately so the columns exist on
    disk even if this run ends up with nothing else to write."""
    ws = wb[SHEET_NAME]
    existing = {str(c.value).strip() for c in ws[1] if c.value}

    missing_base = [c for c in TARGET_COLUMNS if c not in existing]
    if missing_base:
        ref_col = next((c.column for c in ws[1] if str(c.value).strip() == KEY_COLUMN), None)
        if ref_col is None:
            raise ValueError(f"sheet {ws.title!r} is missing expected column {KEY_COLUMN!r}")
        insert_at = ref_col + 1
        ordered_missing = [c for c in BASE_INSERT_ORDER if c in missing_base]
        ws.insert_cols(insert_at, amount=len(ordered_missing))
        for i, name in enumerate(ordered_missing):
            ws.cell(row=1, column=insert_at + i, value=name)
        wb.save(path)
        logger.info(f"inserted base output column(s) {ordered_missing} at position {insert_at} "
                    f"(right after {KEY_COLUMN!r})")
        existing |= set(ordered_missing)

    if FULL_PATH_COLUMN not in existing:
        file_names_col = next(c.column for c in ws[1] if str(c.value).strip() == 'File Names')
        insert_at = file_names_col + 1
        ws.insert_cols(insert_at)
        ws.cell(row=1, column=insert_at, value=FULL_PATH_COLUMN)
        wb.save(path)
        logger.info(f"inserted '{FULL_PATH_COLUMN}' column at position {insert_at} (right after File Names)")


def ensure_supporting_doc_columns(wb, path):
    """Idempotently append the Supporting Document sheet's output columns
    if the sheet exists and doesn't already have them. No-ops (with a
    log) if the workbook has no 'Supporting Document' sheet at all, so
    older/simpler templates aren't broken by this."""
    if SUPPORTING_SHEET_NAME not in wb.sheetnames:
        logger.info(f'workbook has no {SUPPORTING_SHEET_NAME!r} sheet -- skipping supporting-document columns')
        return
    ws = wb[SUPPORTING_SHEET_NAME]
    existing = {str(c.value).strip() for c in ws[1] if c.value}
    missing = [c for c in SUPPORTING_TARGET_COLUMNS if c not in existing]
    if not missing:
        return
    next_col = ws.max_column + 1
    for i, name in enumerate(missing):
        ws.cell(row=1, column=next_col + i, value=name)
    wb.save(path)
    logger.info(f'added supporting-document column(s) {missing} to {SUPPORTING_SHEET_NAME!r}')


def get_pending_rows(wb):
    """Return a list of dicts, one per row that has a Record Reference Id
    but is missing at least one of the target columns (including File
    Path):
        {'row': <1-based row idx>, 'record_reference_id': str}
    Rows with all target columns already filled are skipped (checkpoint).
    Rows with no Record Reference Id at all are blank/padding rows and
    are skipped too.
    """
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f'workbook has no {SHEET_NAME!r} sheet (found: {wb.sheetnames})')
    ws = wb[SHEET_NAME]
    headers = _header_map(ws)

    pending = []
    for row_idx in range(2, ws.max_row + 1):
        key_val = ws.cell(row=row_idx, column=headers[KEY_COLUMN]).value
        if _is_blank(key_val):
            continue
        values = {
            col: ws.cell(row=row_idx, column=headers[col]).value
            for col in ALL_TARGET_COLUMNS
        }
        if all(not _is_blank(v) for v in values.values()):
            continue  # already fully populated -> checkpoint skip
        pending.append({'row': row_idx, 'record_reference_id': str(key_val).strip()})
    return pending


def is_fully_populated(wb):
    return len(get_pending_rows(wb)) == 0


def write_row(wb, path, row_idx, request_id=None, record_id=None, file_names=None, file_path=None):
    """Write whichever of the four target values are provided into the
    given row and save immediately -- one save per record, so a crash
    mid-run loses at most the in-flight record (per spec).

    Only ever touches these columns, and only a cell that is currently
    blank -- never overwrites a value that's already there. This matters
    even within one "pending" row: a row only needs ONE of its target
    cells empty to be picked up at all (see get_pending_rows), so a
    partially-filled row (e.g. Request ID already entered, others still
    blank) must have that existing Request ID left untouched, not
    silently replaced by a freshly-scraped value.
    """
    ws = wb[SHEET_NAME]
    headers = _header_map(ws)
    for col_name, value in (
        ('Request ID', request_id),
        ('Record ID', record_id),
        ('File Names', file_names),
        (FULL_PATH_COLUMN, file_path),
    ):
        if value is None:
            continue
        cell = ws.cell(row=row_idx, column=headers[col_name])
        if _is_blank(cell.value):
            cell.value = value
        else:
            logger.info(f'row {row_idx}: {col_name!r} already has a value ({cell.value!r}) -> leaving it, not overwriting with {value!r}')
    wb.save(path)


def _supporting_header_map(ws):
    headers = {}
    for cell in ws[1]:
        if cell.value:
            headers[str(cell.value).strip()] = cell.column
    missing = [c for c in SUPPORTING_KEY_COLUMNS if c not in headers]
    if missing:
        raise ValueError(f'sheet {ws.title!r} is missing expected column(s): {missing}')
    return headers


def build_supporting_doc_index(wb):
    """Return {(record_reference_id, document_reference_id): row_idx} for
    every row in the Supporting Document sheet, or {} if the workbook has
    no such sheet -- lets callers look up which row to write a scraped
    supporting-document entry into."""
    if SUPPORTING_SHEET_NAME not in wb.sheetnames:
        return {}
    ws = wb[SUPPORTING_SHEET_NAME]
    headers = _supporting_header_map(ws)
    index = {}
    for row_idx in range(2, ws.max_row + 1):
        ref_id = ws.cell(row=row_idx, column=headers['Record Reference Id']).value
        doc_ref = ws.cell(row=row_idx, column=headers['Document Reference Id']).value
        if _is_blank(ref_id) or _is_blank(doc_ref):
            continue
        index[(str(ref_id).strip(), str(doc_ref).strip())] = row_idx
    return index


def write_supporting_doc(wb, row_idx, doc_id=None, filename=None, file_path=None,
                          status=None, error_message=None, parent_request_id=None,
                          parent_record_id=None):
    """Same non-overwrite/blank-cell-only discipline as write_row(), but
    for one row of the Supporting Document sheet. Deliberately does NOT
    save -- a record can have a dozen-plus supporting documents, so the
    caller writes all of them for one record and saves once (alongside
    that record's own write_row call) rather than once per document."""
    ws = wb[SUPPORTING_SHEET_NAME]
    headers = _supporting_header_map(ws)
    for col_name, value in (
        ('Supporting Document ID', doc_id),
        ('Filename', filename),
        ('File Path', file_path),
        ('Status', status),
        ('Error Message', error_message),
        ('Parent Request ID', parent_request_id),
        ('Parent Record ID', parent_record_id),
    ):
        if value is None or col_name not in headers:
            continue
        cell = ws.cell(row=row_idx, column=headers[col_name])
        if _is_blank(cell.value):
            cell.value = value
        else:
            logger.info(f'supporting doc row {row_idx}: {col_name!r} already has a value ({cell.value!r}) -> leaving it')
