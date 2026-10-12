"""Name check (Hound) and Users Collector paging tests."""
from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

import app as app_module  # noqa: E402
from name_check import users_list  # noqa: E402
from name_check import workbook as check  # noqa: E402
from name_check.matcher import NameMatcher  # noqa: E402
from name_check.users_match import MATCHED, NOT_FOUND, REVIEW, User, UserMatcher  # noqa: E402

USERS = [
    User("Sudipta Sen", "s.sen@example.com"),
    User("Chandni Gurnani", "c.gurnani@example.com"),
    User("Niranjan Manohar Lad", "n.lad@example.com"),
    User("Mitika Jawharkar", "m.j@example.com"),
    User("P K Kedhari Sundaram", "pk@example.com"),
    User("Amit Chhatpar", "amit@example.com"),
    User("Rahul Verma", "rahul.old@example.com", active=False),
    User("Rahul Varma", "rahul@example.com"),
    User("Mohammed Ali", "m.ali1@example.com"),
    User("Mohammed Ali", "m.ali2@example.com"),
    User("Ali Syed", "ali.syed@example.com"),
    User("Syed Ali", "syed.ali@example.com"),
    User("Zara Old", "zara@example.com", active=False),
]


@pytest.fixture
def matcher() -> UserMatcher:
    return UserMatcher(USERS)


@pytest.mark.parametrize(
    "query, expected",
    [
        ("Sen, Sudipta", "Sudipta Sen"),  # comma-inverted
        ("Chandni Gurnani, PICG", "Chandni Gurnani"),  # department suffix
        ("Lad, Niranjan Manohar", "Niranjan Manohar Lad"),
        ("Mitika Sharad Jawharkar", "Mitika Jawharkar"),  # extra middle name
        ("P.K. Kedhari Sundaram", "P K Kedhari Sundaram"),  # dotted initials
        ("Amit  Chhatpar .", "Amit Chhatpar"),  # double space, trailing dot
    ],
)
def test_hound_name_variants_still_match(matcher, query, expected):
    m = matcher.match(query)
    assert (m.name, m.status) == (expected, MATCHED), m


def test_scores_agree_with_hound(matcher):
    hound = NameMatcher([u.name for u in USERS])
    for query in ("Sen, Sudipta", "Mitika Sharad Jawharkar", "Amit Chatpar", "Nobody Here"):
        # Scores are rounded to 4 places for display.
        assert matcher.match(query).score == pytest.approx(hound.match(query).score, abs=1e-4)


def test_inactive_users_go_to_review_and_close_active_ones_win(matcher, monkeypatch):
    # Exact name is an inactive account; the active "Rahul Varma" (0.91) is offered instead.
    m = matcher.match("Rahul Verma")
    assert (m.name, m.status) == ("Rahul Verma", REVIEW) and "Inactive user" in m.notes
    assert "Rahul Varma" in [a.name for a in m.alternatives]

    import name_check.users_match as um

    monkeypatch.setattr(um, "ACTIVE_PREFERENCE", 0.1)  # now 0.91 counts as "about the same"
    assert UserMatcher(USERS).match("Rahul Verma").name == "Rahul Varma"

    only_inactive = matcher.match("Zara Old")
    assert (only_inactive.name, only_inactive.status) == ("Zara Old", REVIEW)
    assert "Inactive user" in only_inactive.notes


def test_shared_and_reordered_names_go_to_review(matcher):
    shared = matcher.match("Mohammed Ali")
    assert shared.status == REVIEW and "2 users have this name" in shared.notes
    reordered = matcher.match("Ali Syed")
    assert reordered.status == REVIEW and any("Also close to" in n for n in reordered.notes)


def test_email_and_not_found(matcher):
    m = matcher.match("C.Gurnani@example.com")
    assert (m.name, m.status, m.score) == ("Chandni Gurnani", MATCHED, 1.0)
    nobody = matcher.match("Completely Unknown Person")
    assert (nobody.name, nobody.status) == (None, NOT_FOUND)
    assert matcher.match("   ").status == NOT_FOUND


# ---------- workbook ----------


def hound_workbook() -> bytes:
    """Same layout as Hound's Mashreq template (Sheet3)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet3"
    ws.append(["Business Owner 1", "Is Business Owner 1 Correct / False", "Correct User name ", "Validation Status"])
    ws.append(["Sen, Sudipta", None, None, None])
    ws.append(["Chandni Gurnani", None, None, None])
    ws.append(["Completely Unknown Person", None, None, None])
    ws.append(["Mohammed Ali", None, "Someone Set This", "Matched"])
    other = wb.create_sheet("Ticket Details")
    other.append(["Request ID", "Business Owner 1", "Business Owner 2", "Requesting Department"])
    other.append([1, "Sen, Sudipta", "Amit Chhatpar", "Legal"])
    other.append([2, "Mitika Sharad Jawharkar", None, "IT"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_inspect_lists_sheets_and_suggests_name_columns():
    sheets = {s["name"]: s for s in check.inspect(hound_workbook())["sheets"]}
    assert sheets["Ticket Details"]["suggested"] == ["Business Owner 1", "Business Owner 2"]
    assert sheets["Ticket Details"]["rows"] == 2
    assert sheets["Sheet3"]["suggested"] == ["Business Owner 1"]
    assert sheets["Sheet3"]["samples"]["Business Owner 1"] == "Sen, Sudipta"


def decisions_from(result):
    return [
        {"row": r["row"], "column": r["column"], "name": r["name"], "status": r["status"]}
        for r in result["rows"]
        if r["status"] != check.ALREADY_FILLED
    ]


def test_hound_template_is_filled_in_place(matcher):
    data = hound_workbook()
    result = check.match(data, "Sheet3", ["Business Owner 1"], matcher)
    assert result["stats"] == {MATCHED: 2, REVIEW: 0, NOT_FOUND: 1, check.ALREADY_FILLED: 1}
    assert result["outputs"][0]["correct"] == "Correct User name"

    out = openpyxl.load_workbook(io.BytesIO(check.apply(data, "Sheet3", ["Business Owner 1"], decisions_from(result))))
    ws = out["Sheet3"]
    assert [c.value for c in ws[1]] == [
        "Business Owner 1", "Is Business Owner 1 Correct / False", "Correct User name ", "Validation Status",
    ]
    assert [ws.cell(2, c).value for c in (2, 3, 4)] == ["False", "Sudipta Sen", MATCHED]
    assert [ws.cell(3, c).value for c in (2, 3, 4)] == ["True", "Chandni Gurnani", MATCHED]
    assert [ws.cell(4, c).value for c in (2, 3, 4)] == [None, None, NOT_FOUND]
    assert ws.cell(5, 3).value == "Someone Set This"  # already filled: untouched


def test_several_columns_get_their_own_result_columns(matcher):
    data = hound_workbook()
    cols = ["Business Owner 1", "Business Owner 2"]
    result = check.match(data, "Ticket Details", cols, matcher)
    assert result["stats"][MATCHED] == 3
    ws = openpyxl.load_workbook(io.BytesIO(check.apply(data, "Ticket Details", cols, decisions_from(result))))["Ticket Details"]
    assert [c.value for c in ws[1]] == [
        "Request ID",
        "Business Owner 1", "Business Owner 1 - Correct User name", "Business Owner 1 - Validation Status",
        "Business Owner 2", "Business Owner 2 - Correct User name", "Business Owner 2 - Validation Status",
        "Requesting Department",
    ]
    assert [c.value for c in ws[2]] == [
        1, "Sen, Sudipta", "Sudipta Sen", MATCHED, "Amit Chhatpar", "Amit Chhatpar", MATCHED, "Legal",
    ]
    assert ws.cell(3, 3).value == "Mitika Jawharkar" and ws.cell(3, 6).value is None

    # Checking again sees its own results as already filled.
    again = check.match(check.apply(data, "Ticket Details", cols, decisions_from(result)), "Ticket Details", cols, matcher)
    assert again["stats"][check.ALREADY_FILLED] == 3


def test_apply_rejects_bad_input(matcher):
    data = hound_workbook()
    with pytest.raises(check.CheckError, match="no sheet"):
        check.match(data, "Nope", ["Business Owner 1"], matcher)
    with pytest.raises(check.CheckError, match="not on the sheet"):
        check.match(data, "Sheet3", ["Owner 9"], matcher)
    with pytest.raises(check.CheckError, match="Invalid decision"):
        check.apply(data, "Sheet3", ["Business Owner 1"], [{"row": 2, "column": "Business Owner 1", "status": "Bogus"}])


def test_name_check_api(monkeypatch, tmp_path):
    data_dir = tmp_path / "data"
    monkeypatch.setattr(users_list, "DATA_DIR", data_dir)
    monkeypatch.setattr(users_list, "FILE", data_dir / "users.xlsx")
    monkeypatch.setattr(users_list, "META", data_dir / "meta.json")
    client = TestClient(app_module.app)
    files = {"workbook": ("Owners.xlsx", hound_workbook(), "application/octet-stream")}
    opts = {"options": json.dumps({"sheet": "Sheet3", "columns": ["Business Owner 1"]})}

    no_list = client.post("/api/name-check/match", files=files, data=opts)
    assert no_list.status_code == 400 and "no users list" in no_list.json()["detail"]

    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Email", "Is Active?"])
    for u in USERS:
        wb.active.append([u.name, u.email, "Yes" if u.active else "No"])
    buf = io.BytesIO()
    wb.save(buf)
    assert client.post("/api/users-list", files={"file": ("users.xlsx", buf.getvalue(), "x")}).status_code == 200

    assert client.post("/api/name-check/inspect", files=files).json()["sheets"][0]["name"] == "Sheet3"
    result = client.post("/api/name-check/match", files=files, data=opts).json()
    assert result["stats"][MATCHED] == 2
    applied = client.post(
        "/api/name-check/apply",
        files=files,
        data={"options": json.dumps({"sheet": "Sheet3", "columns": ["Business Owner 1"], "decisions": decisions_from(result)})},
    )
    assert applied.status_code == 200
    assert 'filename="Owners_names-checked.xlsx"' in applied.headers["content-disposition"]
    assert openpyxl.load_workbook(io.BytesIO(applied.content))["Sheet3"].cell(2, 3).value == "Sudipta Sen"


# ---------- Users Collector: verified paging ----------


def load_users_scraper():
    """Import users_collector/scraper.py without clashing with request_id's modules."""
    saved = {k: sys.modules.pop(k) for k in ("config", "scraper") if k in sys.modules}
    path = str(BACKEND / "users_collector")
    sys.path.insert(0, path)
    try:
        import scraper
    finally:
        sys.path.remove(path)
        for k in ("config", "scraper"):
            sys.modules.pop(k, None)
        sys.modules.update(saved)
    return scraper


class FakeSite:
    """Users table with "Showing X to Y of Z" paging; each way of paging can be broken."""

    def __init__(self, total=250, per_page=100, next_works=True, page_buttons_work=True, url_works=True):
        self.total, self.per_page, self.first = total, per_page, 1
        self.next_works, self.page_buttons_work, self.url_works = next_works, page_buttons_work, url_works

    # selenium surface used by the scraper
    def execute_script(self, script, *args):
        if "innerText" in script:
            last = min(self.first + self.per_page - 1, self.total)
            return f"PER PAGE 10 50 100\nShowing {self.first} to {last} of {self.total} results"
        return None

    def find_elements(self, _by, selector):
        if selector == "ul.pagination button":
            return [Button(self, "Go to next page")] + [
                Button(self, f"Go to page {n}") for n in range(1, -(-self.total // self.per_page) + 1)
            ]
        if selector == "table tbody tr":
            return [Row()]
        return []

    def get(self, url):
        page = int(re.search(r"page=(\d+)", url).group(1))
        self.first = (page - 1) * self.per_page + 1 if self.url_works else 1


class Row:
    text = "user row"


class Button:
    def __init__(self, site, label):
        self.site, self.label = site, label

    def get_attribute(self, name):
        return self.label if name == "aria-label" else None

    def is_displayed(self):
        return True

    def is_enabled(self):
        return True

    def click(self):
        s = self.site
        if self.label == "Go to next page" and s.next_works:
            s.first += s.per_page
        elif self.label.startswith("Go to page ") and s.page_buttons_work:
            s.first = (int(self.label.rsplit(" ", 1)[1]) - 1) * s.per_page + 1


@pytest.fixture
def users_scraper(monkeypatch):
    scraper = load_users_scraper()
    monkeypatch.setattr(scraper.time, "sleep", lambda _s: None)
    monkeypatch.setattr(scraper.config, "PAGE_LOAD_WAIT", 0)
    monkeypatch.setattr(scraper.config, "BASE_APP_URL", "https://clm.example.com/t")
    return scraper


def test_parse_showing(users_scraper):
    assert users_scraper.parse_showing("x Showing 101 to 200 of 11206 results y") == (101, 200, 11206)
    assert users_scraper.parse_showing("nothing here") is None


@pytest.mark.parametrize(
    "site",
    [
        FakeSite(),
        FakeSite(next_works=False),
        FakeSite(next_works=False, page_buttons_work=False),
    ],
    ids=["next button", "page button fallback", "url fallback"],
)
def test_goto_page_lands_on_the_right_rows(users_scraper, site):
    assert users_scraper.goto_page_verified(site, 2, 100, timeout=1) == (101, 200, 250)
    assert users_scraper.goto_page_verified(site, 3, 100, timeout=1) == (201, 250, 250)


def test_goto_page_fails_loudly_instead_of_rereading_page_one(users_scraper):
    site = FakeSite(next_works=False, page_buttons_work=False, url_works=False)
    with pytest.raises(RuntimeError, match="could not load page 2"):
        users_scraper.goto_page_verified(site, 2, 100, timeout=1)
