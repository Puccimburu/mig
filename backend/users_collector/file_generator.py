

import os
import shutil
import logging
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

import config

logger = logging.getLogger(__name__)

XLSX_FILENAME = 'users.xlsx'


def new_run_paths():
    """Create output/<timestamp>/ and reset output/latest/."""
    run_id   = datetime.now().strftime('%Y%m%d_%H%M%S')
    run_dir  = os.path.join(config.OUTPUT_DIR, run_id)
    latest_dir = os.path.join(config.OUTPUT_DIR, 'latest')
    os.makedirs(run_dir, exist_ok=True)
    if os.path.isdir(latest_dir):
        for name in os.listdir(latest_dir):
            p = os.path.join(latest_dir, name)
            if os.path.isfile(p):
                os.remove(p)
    else:
        os.makedirs(latest_dir, exist_ok=True)
    return run_id, run_dir, latest_dir


def setup_logging(run_id):
    """Console + per-run log file, both at INFO.

    undetected_chromedriver resets its own logger to match the root level
    on import -- keeping root at INFO avoids the flood of per-HTTP-command
    debug lines that make logs unreadable (confirmed: DEBUG produces ~2MB
    log files). Set VERBOSE_SELENIUM_LOGS=true to opt in when diagnosing
    a driver/network issue.
    """
    log_path = os.path.join(config.LOGS_DIR, f'{run_id}.log')
    verbose  = os.environ.get('VERBOSE_SELENIUM_LOGS', 'false').lower() == 'true'
    level    = logging.DEBUG if verbose else logging.INFO

    root = logging.getLogger()
    root.setLevel(level)
    fmt  = logging.Formatter('%(asctime)s  %(levelname)-8s  %(name)s  %(message)s',
                              datefmt='%Y-%m-%d %H:%M:%S')

    fh = logging.FileHandler(log_path, encoding='utf-8')
    fh.setLevel(level)
    fh.setFormatter(fmt)
    root.addHandler(fh)

    ch = logging.StreamHandler()
    ch.setLevel(logging.INFO)
    ch.setFormatter(fmt)
    root.addHandler(ch)

    return log_path


class XlsxWriter:
    """Appends rows to an xlsx workbook and saves after every page --
    one save per ~100 rows so a crash loses at most one page of users.
    Also keeps output/latest/users.xlsx in sync after each save."""

    _HEADER_FILL  = PatternFill('solid', fgColor='D46A2C')   # Mashreq orange
    _HEADER_FONT  = Font(bold=True, color='FFFFFF', size=11)
    _HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
    _ROW_ALIGN    = Alignment(vertical='center', wrap_text=False)

    def __init__(self, run_dir, latest_dir):
        self.run_path    = os.path.join(run_dir, XLSX_FILENAME)
        self.latest_path = os.path.join(latest_dir, XLSX_FILENAME)
        self._wb = openpyxl.Workbook()
        self._ws = self._wb.active
        self._ws.title = 'Users'
        self._row = 1
        self._write_header()

    def _write_header(self):
        for col_idx, name in enumerate(config.USER_DATA_COLUMNS, start=1):
            cell = self._ws.cell(row=1, column=col_idx, value=name)
            cell.font      = self._HEADER_FONT
            cell.fill      = self._HEADER_FILL
            cell.alignment = self._HEADER_ALIGN
        self._ws.row_dimensions[1].height = 20
        self._row = 2

    def add_page(self, records):
        """Append a page's worth of user records and save immediately."""
        for rec in records:
            for col_idx, key in enumerate(config.USER_DATA_COLUMNS, start=1):
                cell = self._ws.cell(row=self._row, column=col_idx, value=rec.get(key, ''))
                cell.alignment = self._ROW_ALIGN
            self._row += 1
        self._save()

    def _save(self):
        self._wb.save(self.run_path)
        shutil.copyfile(self.run_path, self.latest_path)

    def close(self):
        """Final save and auto-fit column widths."""
        col_widths = {i: len(h) for i, h in enumerate(config.USER_DATA_COLUMNS, start=1)}
        for row in self._ws.iter_rows(min_row=2, max_row=self._ws.max_row):
            for cell in row:
                if cell.value:
                    col_widths[cell.column] = max(col_widths.get(cell.column, 0),
                                                   min(len(str(cell.value)), 60))
        col_letters = {i: openpyxl.utils.get_column_letter(i) for i in col_widths}
        for i, w in col_widths.items():
            self._ws.column_dimensions[col_letters[i]].width = w + 2
        self._save()
        logger.info(f'xlsx saved: {self.run_path} ({self._row - 2} data rows)')
