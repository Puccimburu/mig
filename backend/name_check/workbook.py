"""
Check user-name columns of a client workbook against the CLM users list, and
write the results back (Hound's layout, generalized to any sheet and columns).

For each checked column C the results go into:
- "Correct User name" / "Validation Status" when only one column is checked and the
  sheet already has those columns (Hound's template), plus "Is C Correct / False"
  when present;
- otherwise "C - Correct User name" / "C - Validation Status", inserted right after C
  when missing.

Rows whose correct-name cell is already filled are left alone.
"""
from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .users_match import MATCHED, NOT_FOUND, REVIEW, UserMatcher

ALREADY_FILLED = "Already filled"
STATUSES = (MATCHED, REVIEW, NOT_FOUND)

# Same colours as Hound.
_STYLE = {
    MATCHED: (PatternFill("solid", fgColor="C6EFCE"), Font(color="276221")),
    REVIEW: (PatternFill("solid", fgColor="FFEB9C"), Font(color="9C6500")),
    NOT_FOUND: (PatternFill("solid", fgColor="FFC7CE"), Font(color="9C0006")),
}

_NAME_HINT = re.compile(
    r"owner|assignee|user|requester|requestor|approver|signator|legal party|contact|manager|reviewer",
    re.I,
)
_NOT_NAME = re.compile(r"correct|status|is .* correct|id$|date|department|email", re.I)


class CheckError(Exception):
    pass


def _key(text: object) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def _load(data: bytes, keep_formulas: bool = False):
    try:
        return openpyxl.load_workbook(io.BytesIO(data), data_only=not keep_formulas)
    except Exception as e:
        raise CheckError(f"Not a readable Excel file ({e}).") from e


def _headers(ws: Worksheet) -> dict[str, int]:
    """Normalized header text -> column index (row 1)."""
    return {_key(c.value): c.column for c in ws[1] if c.value is not None and _key(c.value)}


def _sheet(wb, name: str) -> Worksheet:
    if name not in wb.sheetnames:
        raise CheckError(f'The workbook has no sheet "{name}".')
    return wb[name]


def inspect(data: bytes) -> dict[str, Any]:
    wb = _load(data)
    sheets = []
    for ws in wb.worksheets:
        headers = [str(c.value).strip() for c in ws[1] if c.value is not None and str(c.value).strip()]
        samples: dict[str, str] = {}
        cols = _headers(ws)
        for h in headers:
            col = cols[_key(h)]
            for r in range(2, min(ws.max_row, 200) + 1):
                v = ws.cell(r, col).value
                if v is not None and str(v).strip():
                    samples[h] = str(v).strip()
                    break
        rows = sum(
            1 for r in range(2, ws.max_row + 1)
            if any(ws.cell(r, c).value not in (None, "") for c in cols.values())
        )
        sheets.append({
            "name": ws.title,
            "rows": rows,
            "headers": headers,
            "samples": samples,
            "suggested": [h for h in headers if _NAME_HINT.search(h) and not _NOT_NAME.search(h)],
        })
    return {"sheets": sheets}


@dataclass
class Outputs:
    column: str
    correct: str
    status: str
    flag: str | None


def outputs(headers: dict[str, int], columns: list[str]) -> list[Outputs]:
    hound = (
        len(columns) == 1
        and "correct user name" in headers
        and "validation status" in headers
    )
    out = []
    for c in columns:
        if _key(c) not in headers:
            raise CheckError(f'Column "{c}" is not on the sheet.')
        if hound:
            flag = f"is {_key(c)} correct / false"
            out.append(Outputs(c, "Correct User name", "Validation Status",
                               flag if flag in headers else None))
        else:
            out.append(Outputs(c, f"{c} - Correct User name", f"{c} - Validation Status", None))
    return out


def match(data: bytes, sheet: str, columns: list[str], matcher: UserMatcher) -> dict[str, Any]:
    if not columns:
        raise CheckError("Choose at least one name column to check.")
    ws = _sheet(_load(data), sheet)
    headers = _headers(ws)
    rows = []
    stats = {MATCHED: 0, REVIEW: 0, NOT_FOUND: 0, ALREADY_FILLED: 0}
    for o in outputs(headers, columns):
        col = headers[_key(o.column)]
        correct_col = headers.get(_key(o.correct))
        for r in range(2, ws.max_row + 1):
            value = ws.cell(r, col).value
            if value is None or not str(value).strip():
                continue
            current = ws.cell(r, correct_col).value if correct_col else None
            if current is not None and str(current).strip():
                stats[ALREADY_FILLED] += 1
                rows.append({"row": r, "column": o.column, "value": str(value).strip(),
                             "status": ALREADY_FILLED, "name": str(current).strip(), "score": None,
                             "notes": [], "alternatives": [], "email": None})
                continue
            m = matcher.match(value)
            stats[m.status] += 1
            rows.append({
                "row": r,
                "column": o.column,
                "value": m.query,
                "status": m.status,
                "name": m.name,
                "score": m.score,
                "notes": m.notes,
                "alternatives": [asdict(a) for a in m.alternatives],
                "email": m.email,
            })
    return {
        "sheet": sheet,
        "outputs": [asdict(o) for o in outputs(headers, columns)],
        "stats": stats,
        "rows": rows,
    }


def apply(data: bytes, sheet: str, columns: list[str], decisions: list[dict[str, Any]]) -> bytes:
    """Write the chosen names and statuses; returns the updated workbook."""
    wb = _load(data, keep_formulas=True)
    ws = _sheet(wb, sheet)
    plan = outputs(_headers(ws), columns)

    # Create missing output columns right after their input column. Right-most first,
    # so inserting never shifts a column that is still to be handled.
    for o in sorted(plan, key=lambda o: _headers(ws)[_key(o.column)], reverse=True):
        headers = _headers(ws)
        missing = [h for h in (o.correct, o.status) if _key(h) not in headers]
        if missing:
            at = headers[_key(o.column)] + 1
            ws.insert_cols(at, len(missing))
            for i, h in enumerate(missing):
                ws.cell(1, at + i, h).font = Font(bold=True)

    headers = _headers(ws)
    by_column = {_key(o.column): o for o in plan}
    for d in decisions:
        o = by_column.get(_key(d.get("column")))
        status = d.get("status")
        row = d.get("row")
        if o is None or status not in STATUSES or not isinstance(row, int) or row < 2:
            raise CheckError(f"Invalid decision: {d}")
        correct = ws.cell(row, headers[_key(o.correct)])
        if correct.value is not None and str(correct.value).strip():
            continue  # never overwrite a name someone already entered
        name = (d.get("name") or "").strip() or None
        if status == NOT_FOUND:
            name = None
        correct.value = name
        status_cell = ws.cell(row, headers[_key(o.status)])
        status_cell.value = status
        fill, font = _STYLE[status]
        for cell in (correct, status_cell):
            cell.fill, cell.font = fill, font
        if o.flag:
            original = ws.cell(row, headers[_key(o.column)]).value
            ws.cell(row, headers[o.flag]).value = (
                None if name is None else "True" if _key(original) == _key(name) else "False"
            )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
