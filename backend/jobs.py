"""Runs Request ID population jobs as child processes and tracks their progress."""
from __future__ import annotations

import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
import uuid
from collections import deque
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

REQUEST_ID_DIR = Path(__file__).resolve().parent / "request_id"
SHEET_NAME = "Contract Request"
KEY_COLUMN = "Record Reference Id"
REQUIRED_SETTINGS = ("url", "username", "password", "uploaded_on", "total_count")
MAX_FINISHED_JOBS = 5

TOTAL_RE = re.compile(r"matched batch report has (\d+) record")


class JobError(Exception):
    """A request the user can fix (bad input); shown as-is in the UI."""


class JobBusy(JobError):
    pass


@dataclass
class Job:
    id: str
    dir: Path
    workbook_name: str
    status: str = "running"  # running | done | failed | cancelled
    started_at: float = field(default_factory=time.time)
    finished_at: float | None = None
    logs: list[dict[str, Any]] = field(default_factory=list)
    counters: dict[str, int] = field(
        default_factory=lambda: {"total": 0, "processed": 0, "notReady": 0, "noMatch": 0, "errors": 0}
    )
    result: dict[str, Any] | None = None
    stderr_tail: deque = field(default_factory=lambda: deque(maxlen=40))
    process: subprocess.Popen | None = None
    cancelled: bool = False

    def output_file(self, name: str) -> Path | None:
        latest = self.dir / "output" / "latest"
        path = latest / name
        return path if path.is_file() else None

    def workbook(self) -> Path | None:
        return self.output_file(self.workbook_name)

    def report(self) -> Path | None:
        return self.output_file("detailed_report.csv")

    def view(self, since: int = 0) -> dict[str, Any]:
        result = None
        if self.result is not None:
            # Only what the UI needs; never echo anything from the settings.
            result = {
                k: self.result.get(k)
                for k in ("success", "processed", "skipped_not_ready", "skipped_no_match",
                          "fully_completed", "errors", "skipped")
            }
        return {
            "id": self.id,
            "status": self.status,
            "workbookName": self.workbook_name,
            "startedAt": self.started_at,
            "finishedAt": self.finished_at,
            "counters": dict(self.counters),
            "logs": self.logs[since:],
            "logCount": len(self.logs),
            "result": result,
            "hasWorkbook": self.workbook() is not None,
            "hasReport": self.report() is not None,
        }


def check_workbook(path: Path) -> None:
    """Fail fast on a file the populator can't use, before any browser starts."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception as e:
        raise JobError(f"Not a readable Excel file ({e}).") from e
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise JobError(f'The workbook has no "{SHEET_NAME}" sheet. Upload the filled template.')
        header = next(wb[SHEET_NAME].iter_rows(min_row=1, max_row=1, values_only=True), ())
        if KEY_COLUMN not in {str(h).strip() for h in header if h is not None}:
            raise JobError(f'The "{SHEET_NAME}" sheet has no "{KEY_COLUMN}" column.')
    finally:
        wb.close()


def _kill_tree(proc: subprocess.Popen) -> None:
    """Stop the job and the Chrome it started (Chrome runs as a child process)."""
    if proc.poll() is not None:
        return
    if os.name == "nt":
        subprocess.run(["taskkill", "/PID", str(proc.pid), "/T", "/F"], capture_output=True)
    else:
        try:
            os.killpg(proc.pid, signal.SIGTERM)
        except ProcessLookupError:
            pass


class JobManager:
    def __init__(self, python: str | None = None, script_dir: Path = REQUEST_ID_DIR):
        self.python = python or sys.executable
        self.script_dir = script_dir
        self.jobs: dict[str, Job] = {}
        self.lock = threading.Lock()

    def get(self, job_id: str) -> Job | None:
        return self.jobs.get(job_id)

    def running(self) -> Job | None:
        return next((j for j in self.jobs.values() if j.status == "running"), None)

    def start(self, filename: str, data: bytes, settings: dict[str, Any]) -> Job:
        missing = [k for k in REQUIRED_SETTINGS if not str(settings.get(k) or "").strip()]
        if missing:
            raise JobError(f"Missing: {', '.join(missing)}.")
        name = Path(filename or "workbook.xlsx").name
        if not name.lower().endswith(".xlsx"):
            raise JobError("Upload the filled template as an .xlsx file.")

        with self.lock:
            # Concurrent logins on one CLM account break the first session.
            if self.running():
                raise JobBusy("A run is already in progress. Wait for it to finish or cancel it.")
            job_dir = Path(tempfile.mkdtemp(prefix="request-id-"))
            (job_dir / "input").mkdir()
            path = job_dir / "input" / name
            path.write_bytes(data)
            try:
                check_workbook(path)
            except JobError:
                shutil.rmtree(job_dir, ignore_errors=True)
                raise
            job = Job(id=uuid.uuid4().hex[:12], dir=job_dir, workbook_name=name)
            self._evict_old()
            self.jobs[job.id] = job

        env = {**os.environ, "PYTHONUNBUFFERED": "1", "PYTHONIOENCODING": "utf-8"}
        kwargs: dict[str, Any] = {}
        if os.name == "nt":
            kwargs["creationflags"] = subprocess.CREATE_NEW_PROCESS_GROUP
        else:
            kwargs["start_new_session"] = True
        proc = subprocess.Popen(
            [self.python, "job.py", str(job_dir)],
            cwd=self.script_dir,
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            env=env,
            encoding="utf-8",
            errors="replace",
            **kwargs,
        )
        job.process = proc
        # Settings (with the password) only ever travel over this pipe.
        assert proc.stdin is not None
        proc.stdin.write(json.dumps(settings))
        proc.stdin.close()
        threading.Thread(target=self._read_stdout, args=(job,), daemon=True).start()
        threading.Thread(target=self._read_stderr, args=(job,), daemon=True).start()
        return job

    def cancel(self, job_id: str) -> Job | None:
        job = self.jobs.get(job_id)
        if job and job.status == "running" and job.process:
            job.cancelled = True
            _kill_tree(job.process)
        return job

    def _evict_old(self) -> None:
        finished = sorted(
            (j for j in self.jobs.values() if j.status != "running"), key=lambda j: j.started_at
        )
        for job in finished[: max(0, len(finished) - MAX_FINISHED_JOBS + 1)]:
            shutil.rmtree(job.dir, ignore_errors=True)
            del self.jobs[job.id]

    def _read_stderr(self, job: Job) -> None:
        assert job.process and job.process.stderr
        for line in job.process.stderr:
            job.stderr_tail.append(line.rstrip())

    def _read_stdout(self, job: Job) -> None:
        proc = job.process
        assert proc and proc.stdout
        for line in proc.stdout:
            line = line.strip()
            if not line:
                continue
            try:
                msg = json.loads(line)
            except ValueError:
                msg = {"type": "log", "level": "INFO", "logger": "stdout", "message": line, "time": time.time()}
            if msg.get("type") == "result":
                msg.pop("type")
                job.result = msg
            elif msg.get("type") == "log":
                msg.pop("type")
                self._count(job, msg)
                job.logs.append(msg)
        proc.wait()
        self._finish(job)

    @staticmethod
    def _count(job: Job, msg: dict[str, Any]) -> None:
        text = str(msg.get("message", ""))
        c = job.counters
        if m := TOTAL_RE.search(text):
            c["total"] = int(m.group(1))
        elif text.startswith("row ") and "-> Request ID=" in text:
            c["processed"] += 1
        elif "not ready, skipping" in text:
            c["notReady"] += 1
        elif "no report record for" in text or "unrecognized status" in text:
            c["noMatch"] += 1
        if msg.get("level") in ("ERROR", "CRITICAL"):
            c["errors"] += 1

    def _finish(self, job: Job) -> None:
        job.finished_at = time.time()
        r = job.result
        if r is not None:
            c = job.counters
            c["processed"] = int(r.get("processed") or c["processed"])
            c["notReady"] = int(r.get("skipped_not_ready") or c["notReady"])
            c["noMatch"] = int(r.get("skipped_no_match") or c["noMatch"])
            c["errors"] = len(r.get("errors") or []) or c["errors"]
        if job.cancelled:
            job.status = "cancelled"
        elif r is not None and r.get("success"):
            job.status = "done"
        else:
            job.status = "failed"
            if r is None:
                tail = "\n".join(list(job.stderr_tail)[-8:]) or f"exit code {job.process.returncode if job.process else '?'}"
                job.result = {"success": False, "errors": [f"The populator stopped unexpectedly: {tail}"]}
